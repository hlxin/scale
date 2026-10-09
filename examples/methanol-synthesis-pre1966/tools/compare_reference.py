"""Compare the graph's catalytic tests with a reference workbook of methanol-synthesis data.

    python compare_reference.py WORKBOOK.xlsx [--graph methanol-synthesis.graph.json] [--out report.json]

The reference (another group's curation, e.g. ``methanol_synthesis_through1965.xlsx``) lists one row
per test: source, location, conditions (temperature, pressure, space velocity) and results (methanol
carbon yield, conversion). For every row this asks whether the graph holds the same numbers for a
test of the same source — the row's yield (or conversion) together with its temperature — and, field
by field, whether each number appears on some test of that source. Numbers are compared in one unit
per field (°C, bar, h-1, %), converting the graph's printed units exactly; nothing in the graph is
changed. A row not found is a lead for re-curation, not a verdict: the reference may have derived or
misread a value, and the report keeps its location and quote so that the source can decide.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from curation import DATA as EX  # noqa: E402   the curated data: the restricted copy of this example

TO_C = {"°C": lambda x: x, "C": lambda x: x, "deg C": lambda x: x, "K": lambda x: x - 273.15, "°F": lambda x: (x - 32) * 5 / 9,
        "F": lambda x: (x - 32) * 5 / 9}
TO_BAR = {"bar": 1.0, "atm": 1.01325, "atmos.": 1.01325, "at": 0.980665, "ata": 0.980665, "kg/cm2": 0.980665, "kg/cm²": 0.980665,
          "psi": 0.0689476, "lb/in2": 0.0689476, "lb/sq in": 0.0689476, "lbs. per sq. in.": 0.0689476, "MPa": 10.0, "kPa": 0.01,
          "mm Hg": 0.00133322, "mmHg": 0.00133322, "cm Hg": 0.0133322}


def numbers(v, conv=None) -> list[float]:
    """The numbers of a value object (value, lower, upper), converted when a converter applies."""
    if not isinstance(v, dict) or v.get("status", "present") != "present":
        return []
    out = [v[k] for k in ("value", "lower", "upper") if isinstance(v.get(k), (int, float))]
    if conv is None:
        return [float(x) for x in out]
    unit = (v.get("unit") or "").strip()
    if isinstance(conv, dict) and unit in conv:
        f = conv[unit]
        return [float(f(x)) if callable(f) else float(x) * f for x in out]
    return []


def test_values(t: dict) -> dict[str, list[float]]:
    p = t["props"]
    res = p.get("results") or {}
    pct = [x for k in ("yield", "conversion", "selectivity") for x in numbers(p.get(k))]
    pct += [x for k, v in res.items() if isinstance(v, dict) and (v.get("unit") or "").strip() == "%" for x in numbers(v)]
    return {"pct": pct, "T": numbers(p.get("temperature"), TO_C), "P": numbers(p.get("pressure"), TO_BAR),
            "SV": numbers(p.get("space_velocity"))}


def parse_conditions(s: str) -> dict:
    out = {}
    if m := re.search(r"(-?\d+(?:\.\d+)?)\s*°C", s or ""):
        out["T"] = float(m.group(1))
    if m := re.search(r"(\d+(?:\.\d+)?)\s*bar", s or ""):
        out["P"] = float(m.group(1))
    if m := re.search(r"GHSV\s*([\d,]+(?:\.\d+)?)", s or ""):
        out["SV"] = float(m.group(1).replace(",", ""))
    return out


def close(a: float, b: float, field: str) -> bool:
    if field == "pct":
        return abs(a - b) <= 0.51
    if field == "T":
        return abs(a - b) <= 2.0
    return abs(a - b) <= 0.03 * max(abs(b), 1e-9)


def compare(workbook: Path, graph: dict) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    src = {r[0]: str(r[6]) for r in list(wb["Sources"].iter_rows(values_only=True))[2:] if r and r[0]}
    rows = [r for r in list(wb["Data"].iter_rows(values_only=True))[2:] if r and r[0]]
    by_source: dict[str, list[dict]] = defaultdict(list)
    for n in graph["nodes"]:
        if n["label"] == "CatalyticTest":
            slug = n["id"].split("/")[1]
            by_source[slug].append({"id": n["id"], "name": n["props"].get("name"), **test_values(n)})
    report, per_source = [], defaultdict(lambda: {"rows": 0, "row_found": 0, "fields": defaultdict(lambda: [0, 0])})
    for r in rows:
        ref, loc, cond, y, conv = r[0], r[2], r[7], r[8], r[10]
        doi = src.get(ref, "")
        slug = doi.lower().replace("/", "-") if doi.startswith("10.") else doi.lower()
        want = {**parse_conditions(cond or "")}
        if isinstance(y, (int, float)):
            want["pct"] = float(y)
        elif isinstance(conv, (int, float)):
            want["pct"] = float(conv)
        tests = by_source.get(slug, [])
        fields = {f: any(any(close(x, v, f) for x in t[f]) for t in tests) for f, v in want.items()}
        row_found = [t["id"] for t in tests if "pct" in want and any(close(x, want["pct"], "pct") for x in t["pct"])
                     and ("T" not in want or any(close(x, want["T"], "T") for x in t["T"]))]
        ps = per_source[slug]
        ps["rows"] += 1
        ps["row_found"] += bool(row_found)
        for f, ok in fields.items():
            ps["fields"][f][0] += ok
            ps["fields"][f][1] += 1
        report.append({"reference": ref, "source": slug, "location": loc, "conditions": cond, "yield": y, "conversion": conv,
                       "found_in_tests": row_found[:3], "fields_found": fields, "tests_of_source": len(tests)})
    summary = {s: {"rows": v["rows"], "rows_found": v["row_found"],
                   "fields": {f: f"{a}/{b}" for f, (a, b) in sorted(v["fields"].items())}} for s, v in sorted(per_source.items())}
    total = sum(v["rows"] for v in per_source.values())
    found = sum(v["row_found"] for v in per_source.values())
    return {"workbook": workbook.name, "rows": total, "rows_found": found, "by_source": summary,
            "missing": [x for x in report if not x["found_in_tests"]]}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--graph", type=Path, default=EX / "methanol-synthesis.graph.json")
    ap.add_argument("--out", type=Path)
    a = ap.parse_args(argv)
    res = compare(a.workbook, json.loads(a.graph.read_text()))
    if a.out:
        a.out.write_text(json.dumps(res, indent=1, ensure_ascii=False) + "\n")
    print(json.dumps({k: res[k] for k in ("rows", "rows_found", "by_source")}, indent=1, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
