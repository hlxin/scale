"""Export the methanol-synthesis tests into another group's curated-data workbook layout.

    python export_curated_table.py TEMPLATE.xlsx OUT.xlsx [--graph methanol-synthesis.graph.json]

The template's main sheet ("Curated Data": STY, family, loadings, supports, promoters, calcination,
BET, H2/CO2, GHSV, catalyst amount, pressure, temperature, DOI) keeps exactly its headers; its data
rows are replaced by one row per methanol-synthesis CatalyticTest of the graph (a physical
experiment of CO or CO2 hydrogenation to methanol, not an equilibrium test).

Each cell is either **as curated** (the printed value, converted exactly to the column's unit),
**derived** (computed from curated values of the same source by a stated rule, with a cell comment
saying how), or "-" with a **gap** saying what is missing. Derivation rules (the "Notes" sheet
lists them all):

* STY [mg MeOH h-1 gcat-1] = methanol produced per hour / catalyst mass. Methanol per hour comes
  from a printed methanol mass or mole rate, a liquid-methanol volume (x 0.7914 g/cm3), a gas
  volume at N.T.P. (/ 22 414 cm3/mol x 32.042 g/mol), or a condensate volume with its printed
  methanol content by volume; a per-volume space-time yield is divided by the catalyst's bulk
  density. Catalyst mass is printed, or a printed volume x a printed bulk density.
* GHSV [cm3 h-1 gcat-1] = gas flow (at N.T.P. as printed) / catalyst mass, or a per-volume space
  velocity / the bulk density.
* Family, supports, promoters follow the other group's convention (Cu, Pd or In2O3 as the family
  where present, else the mixed oxide, e.g. ZnO-Cr2O3; the other oxides as supports, minor or
  non-oxide components as promoters), unless the source states the roles.
* Loadings [wt.%] from a complete printed composition (every component with a fraction on one
  basis: mass, mole or atom), metal wt.% through formula masses (the template's oxide-to-metal rule).
* MW of supports from their formulas; CR of the family metal from Cordero et al. 2008 (the values
  the template uses: Cu 132, Zn 122, Pd 139, In 142 pm).
* H2/CO2 from a printed ratio or from the H2 and CO2 fractions of the feed on one basis.

The template's DOI_Vlookup data belong to the other group and are cleared. Added sheets: "Row
sources" (graph record, source, printed values, derivations and gaps per row), "Gap summary"
(coverage per column against the other group's data) and "Notes".
"""

from __future__ import annotations

import argparse
import collections
import json
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(Path(__file__).resolve().parent))
from curation import DATA as EX  # noqa: E402   the curated data: the restricted copy of this example
MAIN = "Curated Data"
MS = {"ex:msyn/reaction/co-hydrogenation-to-methanol", "ex:msyn/reaction/co2-hydrogenation-to-methanol",
      "ex:msyn/10.1246-bcsj.37.1817/reaction/methanol-synthesis-from-co-and-d2"}                # as curation.METHANOL_SYNTHESIS
DASH = "-"
M_MEOH, RHO_MEOH, V_MOLAR = 32.042, 0.7914, 22414.0          # g/mol; g/cm3 at 20 °C; cm3/mol at N.T.P.
CM3_FT3, G_LB, CM3_GAL = 28316.85, 453.592, 3785.41

ATOMIC = {"H": 1.008, "C": 12.011, "N": 14.007, "O": 15.999, "Na": 22.990, "Mg": 24.305, "Al": 26.982, "Si": 28.085,
          "P": 30.974, "S": 32.06, "Cl": 35.45, "K": 39.098, "Ca": 40.078, "Ti": 47.867, "V": 50.942, "Cr": 51.996,
          "Mn": 54.938, "Fe": 55.845, "Co": 58.933, "Ni": 58.693, "Cu": 63.546, "Zn": 65.38, "Ga": 69.723, "Br": 79.904,
          "Zr": 91.224, "Mo": 95.95, "Ag": 107.868, "Pd": 106.42, "Cd": 112.414, "In": 114.818, "Sn": 118.71, "I": 126.904,
          "Ba": 137.327, "La": 138.905, "Ce": 140.116, "W": 183.84, "Pt": 195.084, "Pb": 207.2, "Bi": 208.98, "Th": 232.038,
          "U": 238.029, "Be": 9.012, "B": 10.81, "Li": 6.94, "F": 18.998, "Sr": 87.62, "As": 74.922, "Sb": 121.76, "Hg": 200.59}
# Covalent radii, pm (Cordero et al., Dalton Trans. 2008, 2832); Mn, Fe, Co low-spin.
CR = {"Cu": 132, "Zn": 122, "Pd": 139, "In": 142, "Cr": 139, "Ni": 124, "Mn": 139, "Fe": 132, "Co": 126, "Ag": 145, "Pt": 136,
      "Mo": 154, "W": 162, "Mg": 141, "Al": 121, "Th": 206, "U": 196, "Ce": 204, "Zr": 175, "Ti": 160, "V": 153, "Cd": 144,
      "Ba": 215, "Ca": 176, "K": 203, "Na": 166, "Si": 111, "La": 207, "Sn": 139, "Pb": 146, "Bi": 148}
MW_NAMED = {"kieselguhr": ("SiO2", 60.08), "diatomaceous earth": ("SiO2", 60.08), "silica": ("SiO2", 60.08), "silica gel": ("SiO2", 60.08),
            "charcoal": ("C", 12.0), "activated carbon": ("C", 12.0), "carbon": ("C", 12.0), "alumina": ("Al2O3", 101.96)}
CARRIERS = {"ZnO", "Al2O3", "Cr2O3", "SiO2", "MgO", "ThO2", "ZrO2", "CeO2", "TiO2", "Ga2O3", "In2O3", "Fe3O4", "CaO", "BaO", "MnO",
            "CrO3", "La2O3", "UO2", "U3O8", "WO3", "MoO3", "V2O5"}

TO_K = {"°C": lambda x: x + 273.15, "C": lambda x: x + 273.15, "K": lambda x: x, "°F": lambda x: (x - 32) * 5 / 9 + 273.15}
TO_MPA = {"MPa": 1, "kPa": 1e-3, "bar": 0.1, "atm": 0.101325, "ata": 0.0980665, "at": 0.0980665, "kg/cm2": 0.0980665,
          "lb/in2": 0.00689475729, "psi": 0.00689475729, "mm Hg": 0.000133322368, "mmHg": 0.000133322368,
          "cm Hg": 0.00133322368, "torr": 0.000133322368}
TO_G = {"g": 1, "grams": 1, "gram": 1, "gr": 1, "mg": 1e-3, "kg": 1e3}
TO_CM3 = {"cc": 1, "cm3": 1, "ml": 1, "mL": 1, "c.c.": 1, "L": 1000, "l": 1000, "liter": 1000, "litre": 1000, "cu. m": 1e6, "m3": 1e6}
TO_H = {"h": 1, "hr": 1, "hrs": 1, "hours": 1, "hour": 1, "min": 1 / 60}
DENS = {"g/ml": 1, "g/cc": 1, "g/cm3": 1, "g/mL": 1, "kg/m3": 1e-3, "kg/L": 1, "kg/l": 1}
# methanol mass per hour, g/h
MASS_RATE = {"g/h": 1, "g/hr": 1, "mg/h": 1e-3, "kg/h": 1e3, "kg/hr": 1e3, "kg/day": 1e3 / 24, "t/day": 1e6 / 24,
             "kmol/h": 1e3 * M_MEOH, "kmol/uur": 1e3 * M_MEOH, "mol/h": M_MEOH}
LIQ_RATE = {"cm3/h": 1, "cc/h": 1, "ml/h": 1, "ml CH3OH/h": 1, "cc liquid methanol/h": 1, "L/h": 1000, "l/h": 1000, "l/hr": 1000}
GAS_RATE = {"ml N.T.P./h": 1}
# methanol mass per hour per cm3 of catalyst, g/(h cm3)
VOL_STY = {"kg/l hr": 1, "kg/L h": 1, "kg methanol per liter catalyst per hour": 1, "kg per liter of catalyst per day": 1 / 24,
           "kg/day per cu.m. of catalyst": 1e3 / 1e6 / 24, "lb CH3OH/(ft3 catalyst 24 h)": G_LB / CM3_FT3 / 24,
           "lb/(ft3·24 h)": G_LB / CM3_FT3 / 24, "lb methanol/(ft3 catalyst 24 h)": G_LB / CM3_FT3 / 24,
           "lb methanol/(ft3 catalyst·24 h)": G_LB / CM3_FT3 / 24,
           "gal/(24 h·ft3 catalyst space)": CM3_GAL * RHO_MEOH / CM3_FT3 / 24,
           "cc pure methanol/h per 100 cc catalyst": RHO_MEOH / 100, "cc methanol/h per 100 cc catalyst": RHO_MEOH / 100}
MASS_STY = {"g/hr per g catalyst": 1, "g/h per g catalyst": 1}
# gas flow, cm3 (N.T.P.) per hour
FLOW = {"L/h": 1000, "l/hr": 1000, "l/h": 1000, "cc/min": 60, "cc/h": 1, "ml/min": 60, "mmol/min": 60 * V_MOLAR / 1000,
        "Nm3/hr": 1e6, "m3/h": 1e6, "cu. m. per hour": 1e6, "kmol/h": V_MOLAR * 1e3, "ft3/day": CM3_FT3 / 24}
FLOW_PER_VOL = {"L/h per 3 cc catalyst": 1000 / 3}
# space velocity, h-1 (gas volume per catalyst volume per hour)
SV = {"1/h": 1, "h-1": 1, "hr-1": 1, "h^-1": 1, "v/v/h": 1, "vol gas/vol catalyst/h": 1, "cc/hr/cc": 1, "vol/(h·vol catalyst)": 1,
      "ft3/(ft3·h)": 1, "ft3 gas/(ft3 catalyst h) at N.T.P.": 1, "ft3 gas/(ft3 catalyst·h) at N.T.P.": 1, "cc/cc/min": 60,
      "L exit gas (std. cond.) per L catalyst per h": 1}
WT = {"wt%", "wt.%", "wt %", "weight %", "% by weight", "percent by weight", "wt. %"}
MASSY = {"g": 1, "kg": 1000, "mg": 1e-3, "parts by weight": 1, "parts": 1}


# ------------------------------------------------------------------------------------------ small helpers
def iv(v: dict | None, conv) -> tuple[float, float] | None:
    """A value object as an interval (lo, hi) in the target unit, or None."""
    if not isinstance(v, dict) or v.get("status") != "present":
        return None
    unit = (v.get("unit") or "").strip()
    f = conv.get(unit) if isinstance(conv, dict) else conv
    if f is None:
        return None
    g = f if callable(f) else (lambda x, k=f: x * k)
    if isinstance(v.get("value"), (int, float)):
        x = g(float(v["value"]))
        return (x, x)
    if isinstance(v.get("lower"), (int, float)) and isinstance(v.get("upper"), (int, float)):
        return (g(float(v["lower"])), g(float(v["upper"])))
    return None


def sig(x: float, n: int = 3) -> float:
    if x == 0:
        return 0.0
    from math import floor, log10
    return round(x, max(0, n - 1 - int(floor(log10(abs(x))))))


def cell(x: tuple[float, float] | None, digits=3, bound: str | None = None, sigfig: int | None = None):
    if x is None:
        return DASH
    rnd = (lambda v: sig(v, sigfig)) if sigfig else (lambda v: round(v, digits))
    lo, hi = (rnd(x[0]), rnd(x[1]))
    if lo == hi:
        return {"upper": f"≤{lo}", "lower": f"≥{lo}"}.get(bound, lo)
    return f"{lo}–{hi}"


def mul(a, b):
    return None if a is None or b is None else (min(a[0] * b[0], a[1] * b[1]), max(a[0] * b[0], a[1] * b[1]))


def div(a, b):
    if a is None or b is None or b[0] <= 0:
        return None
    return (min(a[0] / b[1], a[1] / b[0]), max(a[0] / b[1], a[1] / b[0]))


def printed(v) -> str:
    if not isinstance(v, dict):
        return ""
    if v.get("status") != "present":
        return v.get("status", "")
    return v.get("as_reported") or " ".join(str(x) for x in (v.get("value", f"{v.get('lower')}–{v.get('upper')}"), v.get("unit") or "")).strip()


def formula_atoms(f: str | None) -> dict[str, float] | None:
    """Element counts of a simple formula (ZnO, Cr2O3, Cu(NO3)2, Zn(NO3)2·6H2O), or None."""
    if not f:
        return None
    total: dict[str, float] = collections.Counter()
    for part in f.replace(" ", "").replace("·", ".").split("."):
        m = re.match(r"^(\d*)(.*)$", part)
        mult = float(m.group(1) or 1)

        def parse(s, i=0):
            counts = collections.Counter()
            while i < len(s):
                if s[i] == "(":
                    inner, i = parse(s, i + 1)
                    k = re.match(r"\d+", s[i:])
                    n = int(k.group(0)) if k else 1
                    i += len(k.group(0)) if k else 0
                    for e, c in inner.items():
                        counts[e] += c * n
                elif s[i] == ")":
                    return counts, i + 1
                else:
                    em = re.match(r"([A-Z][a-z]?)(\d*\.?\d*)", s[i:])
                    if not em or em.group(1) not in ATOMIC:
                        raise ValueError(s)
                    counts[em.group(1)] += float(em.group(2) or 1)
                    i += len(em.group(0))
            return counts, i
        try:
            c, _ = parse(m.group(2))
        except (ValueError, IndexError):
            return None
        for e, n in c.items():
            total[e] += n * mult
    return dict(total) or None


def mw(formula: str | None) -> float | None:
    at = formula_atoms(formula)
    return round(sum(ATOMIC[k] * v for k, v in at.items()), 2) if at else None


def metals(at: dict[str, float]) -> list[str]:
    return [k for k in at if k not in ("O", "H", "C", "N", "S", "Cl", "F", "Br", "I", "P")]


# ------------------------------------------------------------------------------------------ per test
class Row(dict):
    def __init__(self):
        super().__init__()
        self.derived: dict[str, str] = {}
        self.gaps: dict[str, str] = {}

    def put(self, col, value, how=None, gap=None):
        self[col] = value
        if value == DASH and gap:
            self.gaps[col] = gap
        elif value != DASH and how:
            self.derived[col] = how


def rows_from(doc: dict) -> list[Row]:
    N = {n["id"]: n for n in doc["nodes"]}
    out, inc = collections.defaultdict(list), collections.defaultdict(list)
    for l in doc["links"]:
        out[l["source"]].append(l)
        inc[l["target"]].append(l)
    src = {i: n["props"] for i, n in N.items() if n["label"] == "Source"}
    rows = []
    for n in sorted(doc["nodes"], key=lambda n: n["id"]):
        if n["label"] != "CatalyticTest":
            continue
        p = n["props"]
        if p.get("experiment_kind") != "physical" or str(p.get("test_kind", "")).startswith("equilibrium"):
            continue
        links = out[n["id"]]
        if not any(l["type"] == "OF_REACTION" and l["target"] in MS for l in links):
            continue
        res = dict(p.get("results") or {})
        cond = dict(p.get("conditions") or {})
        sources = sorted({l["target"] for l in links if l["type"] == "REPORTED_IN" and l["target"] in src})
        s = src[sources[0]] if sources else {}
        use = [l for l in links if l["type"] == "USES_CATALYST"]
        sample = use[0]["target"] if use else None
        up = (use[0]["props"] or {}) if use else {}
        cats = [l["target"] for l in out.get(sample, []) if l["type"] == "HAS_FORMULATION"] if sample else []
        cat = cats[0] if cats else None
        catp = N[cat]["props"] if cat else {}
        comps = [(l["props"] or {}, N[l["target"]]["props"]) for l in out.get(cat, []) if l["type"] == "HAS_COMPONENT"] if cat else []
        feeds = [(l["target"], N[l["target"]]["props"]) for l in links if l["type"] == "FED_WITH"]
        chars = [N[l["source"]]["props"] for l in (inc.get(sample, []) if sample else []) if l["type"] == "CHARACTERIZES"]
        r = Row()

        # ---------------- catalyst mass and bulk density
        mass = iv(up.get("amount"), TO_G)
        vol = iv(up.get("amount"), TO_CM3)
        rho, rho_how = None, None
        for c in chars:
            for k, v in (c.get("results") or {}).items():
                if "bed_density" in k and rho is None and iv(v, DENS):
                    rho, rho_how = iv(v, DENS), f"bed density {printed(v)}"
        for k in ("packed_density", "bulk_density", "apparent_density", "bed_density"):
            v = (catp.get("properties") or {}).get(k)
            if rho is None and iv(v, DENS):
                rho, rho_how = iv(v, DENS), f"{k.replace('_', ' ')} {printed(v)}"
        if mass is None and vol is not None and rho is not None:
            mass = mul(vol, rho)
            r.put("Amount", cell(mass, 3), how=f"catalyst volume {printed(up.get('amount'))} × {rho_how}")
        elif mass is not None:
            r.put("Amount", cell(mass, 4), how=None if (up.get("amount") or {}).get("unit") == "g" else f"{printed(up.get('amount'))} in g")
        else:
            r.put("Amount", DASH, gap=(f"catalyst volume {printed(up.get('amount'))} printed, no bulk density") if vol else
                  "catalyst amount not reported")

        # ---------------- methanol per hour (g/h), per cm3 of catalyst, or per g
        meoh, meoh_how, per_vol, per_mass = None, None, None, None
        cands = [("production_rate", p.get("production_rate")), ("space_time_yield", p.get("space_time_yield"))] + \
                [(k, v) for k, v in res.items()]
        for key, v in cands:
            if not isinstance(v, dict) or v.get("status") != "present":
                continue
            unit = (v.get("unit") or "").strip()
            text = " ".join(str(v.get(x) or "") for x in ("as_reported", "note", "basis")).lower() + " " + key
            is_condensate = any(w in text for w in ("condensate", "crude", "liquid product", "total product", "raw methanol"))
            is_meoh = ("methanol" in key or "ch3oh" in unit.lower() or "methanol" in unit.lower() or "methyl alcohol" in text
                       or "methanol" in text) and not is_condensate and not any(w in key for w in ("ether", "formate", "acetaldehyde"))
            if unit in MASS_STY:
                per_mass, meoh_how = iv(v, MASS_STY), f"{key} {printed(v)}"
                break
            if unit in VOL_STY and (is_meoh or key == "space_time_yield") and not is_condensate:
                per_vol, meoh_how = iv(v, VOL_STY), f"{key} {printed(v)}"
                break
            if not is_meoh:
                continue
            if unit in MASS_RATE:
                meoh, meoh_how = iv(v, MASS_RATE), f"{key} {printed(v)}"
            elif unit in LIQ_RATE:
                meoh, meoh_how = mul(iv(v, LIQ_RATE), (RHO_MEOH, RHO_MEOH)), f"{key} {printed(v)} × 0.7914 g/cm3 (liquid methanol)"
            elif unit in GAS_RATE:
                meoh, meoh_how = mul(iv(v, GAS_RATE), (M_MEOH / V_MOLAR,) * 2), f"{key} {printed(v)} / 22 414 cm3/mol × 32.042 g/mol"
            if meoh:
                break
        if meoh is None and per_vol is None and per_mass is None:          # condensate × methanol content by volume
            cond_v = next(((k, v) for k, v in [("production_rate", p.get("production_rate"))] + list(res.items())
                           if isinstance(v, dict) and v.get("status") == "present" and (v.get("unit") or "") in LIQ_RATE
                           and any(w in " ".join(str(v.get(x) or "") for x in ("as_reported", "note", "basis")).lower() + k
                                   for w in ("condensate", "crude", "liquid"))), None)
            frac = next(((k, v) for k, v in res.items() if isinstance(v, dict) and v.get("status") == "present"
                         and "methanol" in k and "volume" in (v.get("unit") or "")), None)
            if cond_v and frac:
                meoh = mul(mul(iv(cond_v[1], LIQ_RATE), iv({**frac[1], "unit": "x"}, {"x": 0.01})), (RHO_MEOH,) * 2)
                meoh_how = f"{printed(cond_v[1])} condensate × {printed(frac[1])} methanol by volume × 0.7914 g/cm3"
        if per_mass:
            r.put("STY", cell(mul(per_mass, (1000, 1000)), sigfig=3), how=f"{meoh_how} × 1000 mg/g")
        elif meoh and mass:
            r.put("STY", cell(mul(div(meoh, mass), (1000, 1000)), sigfig=3), how=f"{meoh_how} / catalyst {cell(mass, 4)} g")
        elif per_vol and rho:
            r.put("STY", cell(mul(div(per_vol, rho), (1000, 1000)), sigfig=3), how=f"{meoh_how} / {rho_how}")
        else:
            if meoh or per_vol:
                why = f"methanol rate known ({meoh_how}); " + ("catalyst mass unknown" if meoh else "no bulk density to turn per-volume STY into per-gram")
            else:
                prod = printed(p.get("production_rate")) or printed(res.get("production_rate")) or printed(p.get("space_time_yield"))
                why = ("condensate or crude product without a methanol content by volume" if prod and
                       any(w in json.dumps(p.get("production_rate") or res.get("production_rate") or {}).lower() for w in ("condensate", "crude"))
                       else "no methanol rate printed" if not prod else f"production rate not specific to methanol ({prod})")
            r.put("STY", DASH, gap=why)

        # ---------------- GHSV (cm3 N.T.P. per h per g)
        flow = next(((k, v) for k, v in [("flow_rate", p.get("flow_rate"))] + list(cond.items())
                     if isinstance(v, dict) and (v.get("unit") or "").strip() in FLOW and "flow" in k), None)
        flow_v = next(((k, v) for k, v in [("flow_rate", p.get("flow_rate"))] + list(cond.items())
                       if isinstance(v, dict) and (v.get("unit") or "").strip() in FLOW_PER_VOL), None)
        sv = iv(p.get("space_velocity"), SV)
        if flow and mass:
            r.put("GHSV", cell(div(iv(flow[1], FLOW), mass), sigfig=3), how=f"gas flow {printed(flow[1])} (N.T.P.) / catalyst {cell(mass, 4)} g")
        elif flow_v and rho:
            r.put("GHSV", cell(div(iv(flow_v[1], FLOW_PER_VOL), rho), sigfig=3), how=f"gas flow {printed(flow_v[1])} / {rho_how}")
        elif sv and rho:
            r.put("GHSV", cell(div(sv, rho), sigfig=3), how=f"space velocity {printed(p.get('space_velocity'))} / {rho_how}")
        elif sv and vol and mass:
            r.put("GHSV", cell(div(mul(sv, vol), mass), sigfig=3), how="space velocity × catalyst volume / catalyst mass")
        else:
            svp = printed(p.get("space_velocity"))
            fu = (p.get("flow_rate") or {}).get("unit")
            r.put("GHSV", DASH, gap=("space velocity printed per catalyst volume; no bulk density or catalyst mass" if sv else
                                     "space velocity printed without a unit" if svp and svp != "not_reported" else
                                     "gas flow printed as mass (needs the gas composition)" if fu in ("g/min", "g/h", "g/分") else
                                     "gas flow printed, catalyst mass unknown" if flow else "no gas flow or space velocity printed"))

        # ---------------- composition: family, supports, promoters, loadings, MW, CR
        stated = collections.defaultdict(list)
        for lp, cp in comps:
            stated[lp.get("role") or "not_stated"].append((lp, cp))
        parts = [(cp.get("name") or "", cp.get("formula"), formula_atoms(cp.get("formula")), lp) for lp, cp in comps]
        masses, basis_note = None, None
        fr = [(nm, f, at, iv(lp.get("fraction"), lambda x: x), ((lp.get("fraction") or {}).get("unit") or "").strip(), lp.get("fraction_basis"))
              for nm, f, at, lp in parts if lp.get("role") != "precursor"]
        if fr and all(x[3] for x in fr):
            kinds = {(u, b) for *_, u, b in fr}
            if len(kinds) == 1:
                u, b = next(iter(kinds))
                mid = lambda x: (x[0] + x[1]) / 2                  # noqa: E731
                if u in WT or u in MASSY or (u == "%" and b == "mass"):
                    masses = [(nm, f, at, mid(x)) for nm, f, at, x, *_ in fr]
                    basis_note = f"printed mass composition ({u})"
                elif (u in ("mol %", "mol%") or (u in ("%", "1") and b == "mole")) and all(x[2] for x in fr):
                    masses = [(nm, f, at, mid(x) * mw(f)) for nm, f, at, x, *_ in fr]
                    basis_note = f"printed molar composition ({u}) × formula masses"
                elif (u in ("atomic %", "atom per atom Cu") or (u in ("%", "1") and b == "atom")) and all(
                        x[2] and len(metals(x[2])) == 1 for x in fr):
                    masses = [(nm, f, at, mid(x) * mw(f) / at[metals(at)[0]]) for nm, f, at, x, *_ in fr]
                    basis_note = f"printed atomic composition ({u}) × formula mass per metal atom"
        total = sum(m for *_, m in masses) if masses else None

        def metal_wt(el):
            if not masses or not total:
                return None
            return 100 * sum(m * at.get(el, 0) * ATOMIC[el] / mw(f) for _, f, at, m in masses if at and mw(f)) / total

        els = collections.Counter(el for _, f, at, _ in parts for el in metals(at or {}))
        fam, fam_el, fam_how = None, None, None
        act_f = [cp for lp, cp in stated.get("active_component", []) if formula_atoms(cp.get("formula"))]
        if act_f:
            at = formula_atoms(act_f[0].get("formula"))
            fam = act_f[0].get("formula")
            fam_el = (metals(at) or [None])[0]
        else:
            for el, label in (("Cu", "Cu"), ("Pd", "Pd"), ("In", "In2O3")):
                if el in els:
                    fam, fam_el = label, el
                    break
            if fam is None:
                oxides = [f for _, f, at, _ in parts if f and at and "O" in at and len(metals(at)) == 1 and set(at) <= set(metals(at)) | {"O"}]
                if oxides:
                    order = oxides
                    if masses:
                        mf = {f: m for _, f, _, m in masses}
                        order = sorted(oxides, key=lambda f: -mf.get(f, 0))
                    order = sorted(dict.fromkeys(order), key=lambda f: (f != "ZnO",))   # ZnO first, as the template's ZnO-ZrO2
                    fam = "-".join(order[:2])
                    fam_el = metals(formula_atoms(order[0]))[0]
            if fam:
                fam_how = "family by the template's convention (Cu, Pd or In2O3 where present, else the main oxides); roles not stated in the source"
        if fam:
            r.put("Family", fam, how=fam_how)
            r.put("CR", CR.get(fam_el, DASH), how=f"covalent radius of {fam_el} (Cordero et al. 2008)", gap=f"no covalent radius tabulated for {fam_el}")
        else:
            r.put("Family", DASH, gap="catalyst components not reported" if not parts else "components without formulas")
            r.put("CR", DASH, gap="no family metal")
        lw = metal_wt(fam_el) if fam_el else None
        r.put("Metal Loading", round(lw, 2) if lw is not None else DASH, how=f"{fam_el} wt.% from {basis_note}",
              gap="composition not printed" if not any(lp.get("fraction") for *_, lp in parts) else "composition incomplete or on mixed bases")

        mixed_family = bool(fam and "-" in fam)
        others = [(nm, f, at, lp) for nm, f, at, lp in parts
                  if lp.get("role") not in ("precursor", "active_component")
                  and not (fam_el and at and fam_el in metals(at) and fam in ("Cu", "Pd", "In2O3"))]
        if stated.get("support") or stated.get("promoter"):
            sup = [(cp.get("name"), cp.get("formula"), formula_atoms(cp.get("formula")), lp) for lp, cp in stated.get("support", [])]
            pro = [(cp.get("name"), cp.get("formula"), formula_atoms(cp.get("formula")), lp) for lp, cp in stated.get("promoter", [])]
            role_how = None
        else:
            fam_formulas = set((fam or "").split("-")) if mixed_family else set()
            sup = [] if mixed_family else [x for x in others if x[1] in CARRIERS or (x[0] or "").lower() in MW_NAMED]
            pro = [x for x in others if x not in sup and x[1] not in fam_formulas]
            if masses and total and not mixed_family:
                share = {f: 100 * m / total for _, f, _, m in masses}
                minor = [x for x in sup if share.get(x[1], 100) < 5]
                sup = [x for x in sup if x not in minor]
                pro = minor + pro
            role_how = "role by the template's convention (carrier oxides as supports, other components as promoters); not stated in the source"
        for k, (col, mwcol) in enumerate((("S1", "MW1"), ("S2", "MW2"), ("S3", "MW3"))):
            if k < len(sup):
                nm, f, at, lp = sup[k]
                r.put(col, f or nm, how=role_how)
                m = mw(f) if f else None
                named = MW_NAMED.get((nm or "").lower())
                r.put(mwcol, m if m else (named[1] if named else DASH),
                      how=(f"formula mass of {f}" if m else f"{nm} taken as {named[0]}" if named else None), gap=f"no formula for {nm}")
            else:
                r.put(col, DASH, gap="no further support" if sup else ("mixed-oxide family (the template lists no support)" if mixed_family
                                                                        else "no support in the catalyst" if parts else "catalyst components not reported"))
                r.put(mwcol, DASH, gap="no support")
        mws = [r[c] for c in ("MW1", "MW2", "MW3") if isinstance(r.get(c), (int, float))]
        r.put("MWT", round(sum(mws), 2) if mws else DASH, how="sum of the support formula masses (the template's SUM)", gap="no support")
        for k, (col, lcol) in enumerate((("P1", "P1L"), ("P2", "P2L"))):
            if k < len(pro):
                nm, f, at, lp = pro[k]
                el = (metals(at) or [None])[0] if at else None
                r.put(col, el or f or nm, how=role_how)
                w = metal_wt(el) if el else None
                f_ = lp.get("fraction")
                if w is not None:
                    r.put(lcol, round(w, 2), how=f"{el} wt.% from {basis_note}")
                elif isinstance(f_, dict) and ((f_.get("unit") or "").strip() in WT) and iv(f_, lambda x: x):
                    r.put(lcol, cell(iv(f_, lambda x: x), 2))
                else:
                    r.put(lcol, DASH, gap="promoter amount not printed" if not f_ else "composition incomplete or on mixed bases")
            else:
                r.put(col, DASH, gap="no further promoter" if pro else "no promoter in the catalyst" if parts else "catalyst components not reported")
                r.put(lcol, DASH, gap="no promoter")

        # ---------------- calcination, BET
        calc = [N[l["target"]]["props"] for l in out.get(sample, []) if l["type"] == "TREATED_BY"
                and N[l["target"]]["props"].get("treatment_kind") == "calcination"] if sample else []
        if calc:
            r.put("CalcT", cell(iv(calc[0].get("temperature"), TO_K), 2), gap="calcination temperature not printed")
            r.put("CalcH", cell(iv(calc[0].get("duration"), TO_H), 3), gap="calcination duration not printed")
        else:
            r.put("CalcT", DASH, gap="no calcination reported")
            r.put("CalcH", DASH, gap="no calcination reported")
        bet = None
        for c in chars:
            sa = (c.get("results") or {}).get("surface_area")
            if isinstance(sa, dict) and (sa.get("unit") or "").replace("²", "2") in ("m2/g", "m2 g-1"):
                bet = iv(sa, lambda x: x)
        r.put("SBET", cell(bet, 1), gap="no BET surface area measured")

        # ---------------- H2/CO2
        h2co2, how, feed_has, h2co2_note = None, None, set(), None
        for fid, fp in feeds:
            rr = (fp.get("ratios") or {}).get("h2_co2_ratio")
            if rr and iv(rr, lambda x: x):
                h2co2, how = iv(rr, lambda x: x), None
            fx = {}
            for l in out.get(fid, []):
                if l["type"] == "HAS_COMPONENT":
                    sp = l["target"].rsplit("/", 1)[-1]
                    feed_has.add(sp)
                    fq = (l["props"] or {}).get("fraction")
                    if isinstance(fq, dict) and iv(fq, lambda x: x):
                        fx[sp] = ((fq.get("unit") or "").strip(), iv(fq, lambda x: x))
            co_minor = "co" in fx and fx["co"][0] == fx.get("co2", ("", (0, 0)))[0] and fx["co"][1][0] > fx["co2"][1][1] if "co2" in fx else False
            if h2co2 is None and "h2" in fx and "co2" in fx and fx["h2"][0] == fx["co2"][0] and co_minor:
                h2co2_note = f"CO is the main carbon oxide (H2/CO2 would be {cell(div(fx['h2'][1], fx['co2'][1]), sigfig=3)})"
            elif h2co2 is None and "h2" in fx and "co2" in fx and fx["h2"][0] == fx["co2"][0]:
                h2co2 = div(fx["h2"][1], fx["co2"][1])
                how = (f"H2 / CO2 fractions of the feed ({cell(fx['h2'][1], 2)} / {cell(fx['co2'][1], 2)} {fx['co2'][0] or 'no unit'})"
                       + ("; the feed also contains CO" if "co" in fx or "co" in feed_has else ""))
        if h2co2:
            r.put("H2CO2", cell(h2co2, 2), how=how)
        else:
            h2co = [printed(v) for _, fp in feeds for k, v in (fp.get("ratios") or {}).items() if "co" in k and "co2" not in k]
            r.put("H2CO2", DASH, gap=h2co2_note or ("CO + H2 feed without CO2" + (f" (H2/CO {h2co[0]})" if h2co else "")) if "co2" not in feed_has and "co" in feed_has
                  else "CO2 in the feed, fractions not printed" if "co2" in feed_has else "feed composition not printed")

        # ---------------- conditions and reference
        pv, tv = p.get("pressure"), p.get("temperature")
        r.put("P", cell(iv(pv, TO_MPA), 4, (pv or {}).get("bound")),
              gap="gauge pressure (psig), ambient not stated" if (pv or {}).get("unit") == "psig" else "pressure not printed as a number")
        r.put("T", cell(iv(tv, TO_K), 2, (tv or {}).get("bound")), gap="temperature not printed as a number")
        r.put("DOI", f"https://doi.org/{s['doi']}" if s.get("doi") else DASH, gap="source has no DOI (patent or report); see Row sources")

        r.update({"_id": n["id"], "_test": p.get("name"), "_source": sources[0] if sources else "", "_year": s.get("year"),
                  "_title": s.get("title"), "_catalyst": catp.get("name", ""),
                  "_components": "; ".join(f"{cp.get('name')} ({lp.get('role') or 'not_stated'}" + (f", {printed(lp.get('fraction'))}" if lp.get("fraction") else "") + ")"
                                           for lp, cp in comps),
                  "_feed": "; ".join(fp.get("name", "") for _, fp in feeds),
                  "_ratios": "; ".join(f"{k} {printed(v)}" for _, fp in feeds for k, v in (fp.get("ratios") or {}).items()),
                  "_T": printed(tv), "_P": printed(pv), "_SV": printed(p.get("space_velocity")),
                  "_STY": printed(p.get("space_time_yield")) or printed(res.get("space_time_yield")),
                  "_prod": printed(p.get("production_rate")) or printed(res.get("production_rate")),
                  "_conv": printed(p.get("conversion")), "_amount": printed(up.get("amount")),
                  "_locator": "; ".join(sorted({e.get("locator", "") for e in p.get("evidence") or []}))[:200]})
        rows.append(r)
    return rows


COLS = ["STY", "Family", "Metal Loading", "CR", "S1", "MW1", "S2", "MW2", "S3", "MW3", "MWT", "P1", "P1L", "P2", "P2L",
        "CalcT", "CalcH", "SBET", "H2CO2", "GHSV", "Amount", "P", "T", None, "DOI", None]
SOURCE_COLS = [("Row in main sheet", None), ("scale-graph test id", "_id"), ("Test", "_test"), ("Source id", "_source"),
               ("Year", "_year"), ("Title", "_title"), ("Catalyst", "_catalyst"), ("Components (role, amount as printed)", "_components"),
               ("Feed", "_feed"), ("Feed ratios as printed", "_ratios"), ("Temperature as printed", "_T"), ("Pressure as printed", "_P"),
               ("Space velocity as printed", "_SV"), ("Space-time yield as printed", "_STY"), ("Production rate as printed", "_prod"),
               ("Conversion as printed", "_conv"), ("Catalyst amount as printed", "_amount"), ("Evidence locators", "_locator")]
NOTES = [
    ("What this is", "The pre-1966 methanol-synthesis corpus (108 sources, 1923-1965) of the SCALE graph, in the layout of the "
                     "'Curated data_1234 datapoints' workbook of another group (recent literature). One row per methanol-synthesis catalytic "
                     "test (CO or CO2 hydrogenation to methanol, physical experiment, not an equilibrium test). Headers are exactly the template's."),
    ("Three kinds of cell", "As curated (printed value in the column's unit), derived (computed from curated values of the same source; "
                            "the cell has a comment saying how, and 'Row sources' lists every derivation), or '-' (a gap; the cell comment "
                            "and 'Row sources' say what is missing, and 'Gap summary' counts gaps per column against the other group's data)."),
    ("STY [mgMeOH h-1 gcat-1]", "Methanol per hour / catalyst mass × 1000. Methanol per hour: a printed methanol mass or mole rate; a "
                                "liquid-methanol volume × 0.7914 g/cm3 (density at 20 °C); a gas volume at N.T.P. / 22 414 cm3/mol × "
                                "32.042 g/mol; or a condensate volume × its printed methanol content by volume × 0.7914. A per-volume "
                                "space-time yield (kg/L/h, lb/ft3/24 h, gal/ft3/24 h) is divided by the printed bulk density. Catalyst "
                                "mass: printed, or a printed volume × a printed bulk density. Condensate without a printed methanol "
                                "content by volume is not used (a gap)."),
    ("GHSV [cm3 h-1 gcat-1]", "Gas flow (cm3 at N.T.P., as the sources state or imply) per hour / catalyst mass, or a per-volume space "
                              "velocity (h-1) / bulk density. A space velocity based on exit gas (Brown and Galloway) is used as printed."),
    ("Family, supports, promoters", "If the source states roles, they are used. Otherwise the other group's convention: family Cu, Pd or "
                                    "In2O3 where present, else the main oxides joined (e.g. ZnO-Cr2O3, as their ZnO-ZrO2); carrier oxides "
                                    "(ZnO, Al2O3, Cr2O3, SiO2, MgO, ThO2, …) as supports, and minor (< 5 wt.%) or other components as "
                                    "promoters, named by element as in their data. For a mixed-oxide family there is no support, as in "
                                    "their ZnO-ZrO2 rows. These cells are derived (commented)."),
    ("Loadings [wt.%]", "Metal wt.% of the catalyst as prepared (oxide form), from a complete printed composition: every component with "
                        "a fraction on one basis. Mass basis directly; mole basis × formula masses; atomic basis × formula mass per metal "
                        "atom; metal share of each oxide through its formula (the template's 'Oxide to metal loading conv' rule). A printed "
                        "range is taken at its midpoint for this sum."),
    ("MW of supports, Total MW", "Formula masses (standard atomic weights). Kieselguhr and silica are taken as SiO2 (60.08, as the "
                                 "template's SBA-15 = 60), carbon as C (12, as the template's CNT). Total = sum, as the template's SUM."),
    ("CR Metal [pm]", "Covalent radius of the family metal, Cordero et al., Dalton Trans. 2008, 2832 (the template's values: Cu 132, "
                      "Zn 122, Pd 139, In 142)."),
    ("H2/CO2 [-]", "A printed H2/CO2 ratio, or the feed's H2 and CO2 fractions on one basis. Most pre-1966 feeds are CO + H2: then '-' "
                   "with the H2/CO ratio in the gap note."),
    ("Pressure, temperature", "Exact unit conversions: atm × 0.101325, kg/cm2 and at/ata × 0.0980665, lb/in2 and psi × 0.00689476, "
                              "mm Hg × 1.33322e-4, cm Hg × 1.33322e-3 MPa; °C + 273.15 K. Gauge pressures (psig) are '-'."),
    ("Ranges and bounds", "A printed range stays a range ('low–high') and carries through derivations as an interval; a bound is '≤x' "
                          "or '≥x'."),
    ("Calcination, SBET", "Only as reported: a calcination treatment of the tested sample; a BET surface area in m2/g."),
    ("Reference DOI", "The source's DOI; patents and reports have none ('-'), and 'Row sources' identifies them."),
    ("Evidence", "Every row traces to a CatalyticTest record of the SCALE graph (methanol-synthesis.graph.json, GALAXY bundle 3.0.0) "
                 "whose values carry verbatim quotes and page locators."),
]


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("template", type=Path)
    ap.add_argument("out", type=Path)
    ap.add_argument("--graph", type=Path, default=EX / "methanol-synthesis.graph.json")
    a = ap.parse_args(argv)
    rows = rows_from(json.loads(a.graph.read_text()))
    wb = openpyxl.load_workbook(a.template)
    ws = wb[MAIN]
    headers = [c.value for c in ws[1]] + [None] * (len(COLS) - ws.max_column)
    theirs = [dict(zip(headers, [c.value for c in rr])) for rr in openpyxl.load_workbook(a.template, data_only=True)[MAIN].iter_rows(min_row=2)
              if rr[0].value is not None]
    style = [copy(ws.cell(2, c)._style) for c in range(1, len(COLS) + 1)]
    ws.delete_rows(2, ws.max_row)
    for i, row in enumerate(rows, start=2):
        for c, key in enumerate(COLS, start=1):
            x = ws.cell(i, c)
            x._style = copy(style[c - 1])
            if key:
                x.value = row[key]
                if key in row.derived:
                    x.comment = Comment("Derived: " + row.derived[key], "SCALE")
                elif key in row.gaps:
                    x.comment = Comment("Gap: " + row.gaps[key], "SCALE")
    dv = wb["DOI_Vlookup"]                                     # the other group's data: cleared
    dv.delete_rows(2, dv.max_row)
    names = dict(zip(COLS, headers))

    rs = wb.create_sheet("Row sources")
    rs.append([h for h, _ in SOURCE_COLS] + ["Derived cells", "Gaps"])
    for i, row in enumerate(rows, start=2):
        rs.append([i] + [row[k] for _, k in SOURCE_COLS[1:]]
                  + ["; ".join(f"{names[k].strip()}: {v}" for k, v in row.derived.items()),
                     "; ".join(f"{names[k].strip()}: {v}" for k, v in row.gaps.items())])

    gs = wb.create_sheet("Gap summary")
    gs.append(["Column", f"Other group (recent literature): filled, of {len(theirs)}", "Pre-1966: as curated", "Pre-1966: derived",
               f"Pre-1966: gap, of {len(rows)}", "Main reasons for the gaps (rows)"])
    summary = {}
    for key in [k for k in COLS if k]:
        h = names[key]
        th = sum(1 for t in theirs if t.get(h) not in (None, 0, "0", "", DASH))
        der = sum(1 for r in rows if key in r.derived and r[key] != DASH)
        gap = sum(1 for r in rows if r[key] == DASH)
        cur = len(rows) - der - gap
        reasons = collections.Counter(re.sub(r"\s*\((?:H2/CO|printed|lb|kg|g|cc|L)[^)]*\)", "", re.sub(r"\(methanol rate known.*", "(methanol rate known)", r.gaps.get(key, ""))).strip()
                                      for r in rows if r[key] == DASH)
        gs.append([h.strip(), th, cur, der, gap, "; ".join(f"{k} ({v})" for k, v in reasons.most_common(4) if k)])
        summary[h.strip()] = {"theirs": th, "curated": cur, "derived": der, "gap": gap}
    notes = wb.create_sheet("Notes")
    notes.append(["Topic", "How the pre-1966 rows were filled"])
    for t, d in NOTES:
        notes.append([t, d])
    for sh in (rs, gs, notes):
        for c in sh[1]:
            c.font = openpyxl.styles.Font(name="Arial", bold=True)
            c.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.font = openpyxl.styles.Font(name="Arial")
                c.alignment = openpyxl.styles.Alignment(wrap_text=sh is not rs, vertical="top")
    notes.column_dimensions["A"].width, notes.column_dimensions["B"].width = 28, 120
    for col, w in zip("ABCDEF", (30, 18, 16, 16, 16, 100)):
        gs.column_dimensions[col].width = w
    for col, w in zip("ABCDEFGHIJKLMNOPQRST", (10, 48, 40, 26, 6, 40, 36, 60, 30, 30, 18, 18, 22, 26, 22, 14, 16, 22, 80, 80)):
        rs.column_dimensions[col].width = w
    rs.freeze_panes = "C2"
    wb.calculation.fullCalcOnLoad = True                      # the template's helper sheets keep their formulas
    a.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    print(json.dumps({"rows": len(rows), "columns": summary}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
